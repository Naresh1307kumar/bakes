from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Set margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

# Add Name
name = doc.add_heading('Naresh Kumar R.', 0)
name.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add Contact Info
contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.add_run('Phone: 9489771511 | Email: naresk1213@gmail.com | GitHub: naresh kumar (Github_id nar)')

def add_heading(text):
    h = doc.add_heading(text, level=1)
    return h

def add_bullet(text):
    doc.add_paragraph(text, style='List Bullet')

# Career Objective
add_heading('Career Objective')
doc.add_paragraph("To work as a top-notch frontend developer using Java, building efficient applications with minimal bugs based on specific requirements. I aim to consistently maintain application development and implement Java's portability and platform-independent features to create better, user-friendly applications.")

# Academics
add_heading('Academics')
table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Course'
hdr_cells[1].text = 'Institution'
hdr_cells[2].text = 'Board / University'
hdr_cells[3].text = 'YOP'
hdr_cells[4].text = 'Performance'

data = [
    ('BSc', 'TCA', 'Madurai Kamaraj University', '2027', '80% (up to 4th Sem)'),
    ('XII', 'VHN', 'State Board', '2024', '93%'),
    ('X', 'VHN', 'State Board', '2022', '83.6%')
]
for course, inst, board, yop, perf in data:
    row_cells = table.add_row().cells
    row_cells[0].text = course
    row_cells[1].text = inst
    row_cells[2].text = board
    row_cells[3].text = yop
    row_cells[4].text = perf

# Curricular Activities
add_heading('Curricular Activities')
doc.add_paragraph("Active participant in technical and developmental activities to enhance practical knowledge and skills.")

# Internship
add_heading('Internship')
p = doc.add_paragraph()
p.add_run('Frontend Developer Intern | DCI').bold = True
add_bullet('Successfully completed an internship focusing on frontend development.')
add_bullet('Projects Involved: PAWS and Solar Panel.')

# Projects
add_heading('Projects')
p = doc.add_paragraph()
p.add_run('Automatic Solar Panel Tracking System').bold = True
doc.add_paragraph('Developed a system where the solar panel automatically tracks and turns towards the sunlight to maximize energy efficiency.', style='List Bullet 2')

p = doc.add_paragraph()
p.add_run('Secure Chit Fund Management System').bold = True
doc.add_paragraph('Developed using Python and PostgreSQL.', style='List Bullet 2')

p = doc.add_paragraph()
p.add_run('MediSure').bold = True

p = doc.add_paragraph()
p.add_run('PAWS – Rabbit Protection Project').bold = True
doc.add_paragraph('Project related to preventing rabbits from causing damage.', style='List Bullet 2')

# Seminar Attended
add_heading('Seminar Attended')
add_bullet('Attended GDG Seminar at TCA.')
add_bullet('Attended GDG Seminar at HCL.')

# Seminar Presented
add_heading('Seminar Presented')
add_bullet('Presented the basic concepts of Java.')
add_bullet('Presented advanced concepts of C.')

# Paper Presentation
add_heading('Paper Presentation')
add_bullet('Presented a paper on "PAWS" and secured the First Prize.')

# Achievements / Recognition
add_heading('Achievements / Recognition')
add_bullet('Winner in the National Science Expo at TCA.')
add_bullet('Overall Runner-up at ANJAC (Innovation Pitch).')
add_bullet('Winner at TCA (Innovation Pitch).')

# Training Undergone
add_heading('Training Undergone')
add_bullet('16 Hrs of Communication Training.')
add_bullet('30 Hrs of Aptitude Training.')
add_bullet('60 Hrs of Technical Skills Training.')

# Workshops Attended
add_heading('Workshops Attended')
add_bullet('Technical Event.')
add_bullet('Critical Thinking.')
add_bullet('Use of AI Tools.')
add_bullet('Introduction to Machine Learning.')

# Certification
add_heading('Certification')
add_bullet('Certified as a Full-Stack Developer by Beestrive (6 Months).')

# Technical Skills
add_heading('Technical Skills')
add_bullet('Java')
add_bullet('C')
add_bullet('Full-Stack Development')
add_bullet('Project Management')

# Management Skills
add_heading('Management Skills')
add_bullet('Interpersonal Skills')
add_bullet('Leadership Skills')
add_bullet('Time and Stress Management Skills')

# Social Services
add_heading('Social Services')
add_bullet('Active Blood Donor at NSS.')
add_bullet('NSS Volunteer attending camp at Thiruvathavur.')

# Area of Interest
add_heading('Area of Interest')
add_bullet('Networking')
add_bullet('Frontend Development')
add_bullet('Backend Development')

# Personal Information
add_heading('Personal Information')
add_bullet('Date of Birth: 13-12-2006')
add_bullet('Languages Known: Tamil, English')
add_bullet('Marital Status: Unmarried')

# Special Interest
add_heading('Special Interest')
add_bullet('Reading documentation about quantum computers.')

# Declaration
add_heading('Declaration')
doc.add_paragraph('I hereby declare that the information provided above is true to the best of my knowledge.')
import os
output_path = os.path.join(os.path.dirname(__file__), 'Naresh_Resume.docx')
doc.save(output_path)
